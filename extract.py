from docx import Document
from openpyxl import Workbook

# Load the Word document
document = Document('test.docx')

# Create a new Excel workbook
workbook = Workbook()
worksheet = workbook.active

# Loop through each table in the document
for table in document.tables:
    # Loop through each row in the table
    for row in table.rows:
        # Create a new row in the Excel worksheet
        new_row = []
        # Loop through each cell in the row
        for cell in row.cells:
            # Extract the text from the cell
            text = ''
            for paragraph in cell._element.xpath('.//w:p'):
                for run in paragraph.xpath('.//w:r'):
                    text += run.text
            # Append the text from the cell to the new row
            new_row.append(text)
        # Write the new row to the Excel worksheet
        worksheet.append(new_row)

# Save the Excel workbook
workbook.save('tables.xlsx')

